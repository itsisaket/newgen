from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ชุดเครื่องมือเก็บข้อมูลการจัดการสวนทุเรียนรายปี_DRFI.docx"
FONT = "Noto Sans Thai"
GREEN = "1F6B35"
PALE_GREEN = "EAF4EC"
PALE_GRAY = "F5F6F7"
GRAY = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_font(run, size=10.5, bold=False, color=BLACK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for slot in ("ascii", "hAnsi", "eastAsia"):
        rpr.rFonts.set(qn(f"w:{slot}"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    node = tcpr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tcpr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, top=90, start=90, bottom=90, end=90):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_margins = tcpr.first_child_found_in("w:tcMar")
    if tc_margins is None:
        tc_margins = OxmlElement("w:tcMar")
        tcpr.append(tc_margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color=GRAY):
    tblpr = table._tbl.tblPr
    node = tblpr.first_child_found_in("w:tblBorders")
    if node is None:
        node = OxmlElement("w:tblBorders")
        tblpr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = node.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            node.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), color)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def text(doc, value="", size=10.5, bold=False, align=None, before=0, after=4, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    set_font(p.add_run(value), size=size, bold=bold)
    return p


def line(doc, label, count=48, suffix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(label + "  "), bold=True)
    set_font(p.add_run("_" * count))
    if suffix:
        set_font(p.add_run("  " + suffix))
    return p


def options(doc, label, values, suffix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    if label:
        set_font(p.add_run(label + "  "), bold=True)
    for idx, value in enumerate(values):
        set_font(p.add_run(("   " if idx else "") + "☐ " + value))
    if suffix:
        set_font(p.add_run("  " + suffix))
    return p


def section(doc, code, title, intro=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GREEN)
    ppr.append(shd)
    set_font(p.add_run(f"{code}  {title}"), size=13, bold=True)
    if intro:
        text(doc, intro, size=9.5, after=6)


def title(doc, value, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(value), size=18, bold=True)
    if subtitle:
        text(doc, subtitle, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=9)


def table(doc, headers, rows, widths, font_size=8.3, first_col_left=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    repeat_header(t.rows[0])
    for i, value in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(cell, 95, 75, 95, 75)
        shade(cell, GREEN)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), size=font_size, bold=True, color=WHITE)
    for ridx, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cells[i], 100, 75, 100, 75)
            if ridx % 2:
                shade(cells[i], PALE_GRAY)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (first_col_left and i == 0) else WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(value)), size=font_size)
    text(doc, "", size=2, after=1)
    return t


def page(doc, value, subtitle=None):
    doc.add_page_break()
    title(doc, value, subtitle)


def footer(section_obj):
    p = section_obj.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("DRFIS  ชุดเครื่องมือเก็บข้อมูลภาคสนามรายปี  |  หน้า "), size=8, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.5)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.58)
sec.right_margin = Inches(0.58)
footer(sec)

for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(10.5)

# 1 ปก
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(35)
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("ชุดเครื่องมือเก็บข้อมูลการจัดการสวนทุเรียนรายปี"), size=22, bold=True)
text(doc, "สำหรับเก็บข้อมูลภาคสนามและนำเข้าระบบ DRFIS", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
text(doc, "พื้นที่ดำเนินงาน จังหวัดศรีสะเกษ", size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
section(doc, "รหัสชุดข้อมูล", "ข้อมูลกำกับแบบสำรวจ")
line(doc, "รหัสแบบสำรวจ", 22, "รหัสครัวเรือน ____________________")
line(doc, "ปีการผลิต", 18, "วันที่เริ่มรอบ ____ / ____ / ____  วันที่สิ้นสุดรอบ ____ / ____ / ____")
line(doc, "ชื่อผู้ให้ข้อมูล", 34, "เบอร์โทร ____________________")
line(doc, "ผู้เก็บข้อมูล", 34, "หน่วยงาน ____________________")
line(doc, "หมู่บ้าน", 20, "ตำบล __________  อำเภอ __________  จังหวัด ศรีสะเกษ")
options(doc, "สถานะแบบ", ["เริ่มเก็บข้อมูล", "ติดตามระหว่างปี", "ปิดรอบปีแล้ว"])
section(doc, "วัตถุประสงค์", "การใช้ชุดเครื่องมือ")
text(doc, "ใช้บันทึกกิจกรรมสวนตลอดหนึ่งรอบปี ตั้งแต่เริ่มรอบการผลิตจนถึงเก็บเกี่ยวและขาย เพื่อนำเข้าข้อมูลในระบบอย่างครบถ้วน ตรวจสอบย้อนกลับได้ และลดการจำข้อมูลย้อนหลัง", size=10.5, after=8)
options(doc, "ความยินยอม", ["ยินยอมให้บันทึกข้อมูล", "ยินยอมให้บันทึกพิกัด", "ยินยอมให้ถ่ายภาพหลักฐาน"])
line(doc, "ลายมือชื่อผู้ให้ข้อมูล", 27, "วันที่ ____ / ____ / ____")

# 2 คู่มือ
page(doc, "คู่มือการใช้ชุดเครื่องมือ", "อ่านก่อนเริ่มสัมภาษณ์และเก็บข้อมูล")
section(doc, "1", "หลักการเก็บข้อมูล")
for item in [
    "ใช้หนึ่งชุดต่อหนึ่งครัวเรือนต่อหนึ่งรอบปีการผลิต และระบุสวนกับแปลงทุกครั้งที่บันทึกรายการ",
    "บันทึกทันทีที่เกิดกิจกรรม หรือรวบรวมอย่างน้อยเดือนละหนึ่งครั้งจากสมุดสวน ใบเสร็จ ภาพถ่าย หรือคำยืนยันของเกษตรกร",
    "ใช้หน่วยมาตรฐานเดียวกัน เช่น ไร่ กิโลกรัม ลิตร ชั่วโมง บาท และกิโลวัตต์ชั่วโมง",
    "หากไม่ทราบ ให้เขียนว่า ไม่ทราบ ห้ามคาดเดา หากไม่มีรายการให้ใส่ 0 หรือทำเครื่องหมาย ไม่มี",
    "ข้อมูลส่วนบุคคล พิกัด และภาพถ่ายต้องมีความยินยอมก่อนบันทึกในระบบ",
]:
    options(doc, "", [item])
section(doc, "2", "รหัสหลักที่ต้องใช้ตรงกัน")
table(doc, ["รหัส", "ตัวอย่าง", "ใช้เชื่อมข้อมูล"], [
    ["ครัวเรือน", "HH-0013", "สมาชิก กลุ่มเกษตรกร หมู่บ้าน"],
    ["สวน", "FARM-001", "เจ้าของสวน ที่ตั้ง แหล่งน้ำ"],
    ["แปลง", "PLOT-01", "กิจกรรม ปัจจัย ผลผลิต การขาย"],
    ["รอบปีผลิต", "2569/70", "วันเริ่ม วันเก็บผลสุดท้าย สรุปผล"],
], [1.25, 1.45, 4.3], font_size=9, first_col_left=True)
section(doc, "3", "รหัสแหล่งข้อมูล")
options(doc, "ใช้รหัส", ["D สมุดหรือเอกสาร", "R ใบเสร็จ", "P ภาพถ่าย", "G พิกัด", "I สัมภาษณ์", "O สังเกตการณ์"])
line(doc, "ผู้ตรวจสอบความครบถ้วน", 32, "วันที่ ____ / ____ / ____")

# 3 A
page(doc, "แบบฟอร์ม A  ทะเบียนครัวเรือน สวน และแปลง", "กรอกครั้งแรกและแก้ไขเมื่อข้อมูลเปลี่ยน")
section(doc, "A1", "ครัวเรือนและที่ตั้ง")
line(doc, "รหัสครัวเรือน", 22, "ชื่อหัวหน้าครัวเรือน ______________________________")
line(doc, "เบอร์โทร", 22, "กลุ่มเกษตรกร ______________________________")
line(doc, "หมู่บ้าน", 20, "หมู่ที่ ____  ตำบล __________  อำเภอ __________")
options(doc, "สถานะ", ["ใช้งานอยู่", "ระงับชั่วคราว", "ถอนตัว"])
section(doc, "A2", "ทะเบียนสวน")
table(doc, ["ลำดับ", "รหัสสวน", "ชื่อสวนหรือจุดสังเกต", "พื้นที่ไร่", "แหล่งน้ำ", "พิกัด", "สถานะ"],
      [[str(i), "", "", "", "", "", ""] for i in range(1, 5)],
      [0.45, 0.85, 1.8, 0.7, 1.0, 1.25, 0.9], font_size=8.3)
section(doc, "A3", "ทะเบียนแปลง")
table(doc, ["รหัสแปลง", "รหัสสวน", "พื้นที่ไร่", "จำนวนต้น", "พันธุ์หลัก", "ปีปลูก", "ระบบน้ำ", "สถานะ"],
      [["", "", "", "", "", "", "", ""] for _ in range(6)],
      [0.85, 0.75, 0.72, 0.75, 1.0, 0.7, 1.1, 0.9], font_size=8.3)

# 4 B
page(doc, "แบบฟอร์ม B  เปิดรอบปีการผลิต", "กำหนดขอบเขตข้อมูลก่อนเริ่มบันทึกรายการรายเดือน")
section(doc, "B1", "ข้อมูลรอบปี")
line(doc, "ชื่อรอบปีการผลิต", 25, "เช่น 2569/70")
line(doc, "วันที่เริ่มรอบ", 20, "วันที่เก็บผลสุดท้ายของรอบก่อน ____ / ____ / ____")
line(doc, "แปลงที่อยู่ในรอบนี้", 52)
line(doc, "พื้นที่ให้ผลผลิต", 16, "ไร่  จำนวนต้นให้ผล ________ ต้น")
options(doc, "สถานะรอบ", ["กำลังดำเนินการ", "ปิดรอบแล้ว"])
section(doc, "B2", "เป้าหมายก่อนเริ่มรอบ")
line(doc, "ผลผลิตเป้าหมาย", 16, "กิโลกรัม  ราคาขายเป้าหมาย ________ บาทต่อกิโลกรัม")
line(doc, "งบประมาณต้นทุน", 18, "บาท")
line(doc, "ตลาดหรือผู้ซื้อเป้าหมาย", 49)
line(doc, "มาตรฐานหรือเงื่อนไขคุณภาพ", 44)
line(doc, "ความเสี่ยงสำคัญที่คาดไว้", 50)
section(doc, "B3", "ข้อมูลฐานจากรอบก่อน")
table(doc, ["ตัวชี้วัด", "ค่ารอบก่อน", "หน่วย", "แหล่งข้อมูล"], [
    ["ผลผลิตรวม", "", "กิโลกรัม", ""], ["ต้นทุนรวม", "", "บาท", ""],
    ["รายได้รวม", "", "บาท", ""], ["ราคาขายเฉลี่ย", "", "บาทต่อกก.", ""],
    ["ผลผลิตเฉลี่ย", "", "กก.ต่อไร่", ""],
], [2.6, 1.4, 1.35, 1.65], font_size=9, first_col_left=True)

# 5-6 ปฏิทิน
for half, months in (("มกราคมถึงมิถุนายน", ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน"]),
                     ("กรกฎาคมถึงธันวาคม", ["กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"])):
    page(doc, f"แบบฟอร์ม C  ปฏิทินการจัดการสวน {half}", "สรุปกิจกรรมเด่น ปัญหา และหลักฐานประจำเดือน")
    section(doc, "C", "ภาพรวมรายเดือน", "หากเดือนใดไม่มีรายการ ให้เขียนว่า ไม่มี")
    rows = [[m, "", "", "", ""] for m in months]
    table(doc, ["เดือน", "กิจกรรมสำคัญ", "ปัจจัยหรือแรงงาน", "ปัญหาหรือสภาพอากาศ", "หลักฐาน"], rows,
          [1.0, 2.05, 1.5, 1.75, 0.85], font_size=8.7, first_col_left=True)
    section(doc, "C ตรวจสอบ", "รายการที่ควรพบในช่วงครึ่งปี")
    options(doc, "", ["ให้น้ำ", "ใส่ปุ๋ย", "ป้องกันศัตรูพืช", "ตัดแต่งกิ่ง", "กำจัดวัชพืช"])
    options(doc, "", ["ระยะดอกและผล", "เก็บเกี่ยว", "ขาย", "จัดการกิ่ง", "ใช้เทคโนโลยี"])
    line(doc, "เหตุการณ์ผิดปกติหรือภัยพิบัติ", 47)
    line(doc, "หมายเหตุส่งต่อผู้บันทึกระบบ", 49)

# 7-8 D activity logs
for part in (1, 2):
    page(doc, f"แบบฟอร์ม D  บันทึกกิจกรรมสวน รายการต่อเนื่อง {part}", "ใช้หนึ่งบรรทัดต่อหนึ่งกิจกรรมและระบุแปลงทุกครั้ง")
    table(doc, ["วันที่", "แปลง", "กิจกรรม", "วัสดุหรือรายละเอียด", "ปริมาณ", "หน่วย", "ค่าใช้จ่าย", "แรงงานชม.", "หลักฐาน"],
          [["", "", "", "", "", "", "", "", ""] for _ in range(12)],
          [0.65, 0.55, 0.95, 1.25, 0.58, 0.55, 0.72, 0.72, 0.75], font_size=7.8)
    text(doc, "ประเภทกิจกรรมที่แนะนำ  ให้น้ำ ใส่ปุ๋ย พ่นสาร ตัดแต่งกิ่ง กำจัดวัชพืช สำรวจโรคแมลง ซ่อมระบบน้ำ เก็บเกี่ยว และกิจกรรมอื่น", size=9, after=7)
    line(doc, "รวมค่าใช้จ่ายหน้านี้", 17, "บาท  รวมแรงงาน ________ ชั่วโมง")
    line(doc, "ผู้บันทึก", 26, "วันที่ตรวจสอบ ____ / ____ / ____")

# 9 E
page(doc, "แบบฟอร์ม E  ปัจจัยการผลิต น้ำ และพลังงาน", "สรุปยอดทั้งปีจากแบบฟอร์ม D และหลักฐานประกอบ")
section(doc, "E1", "ปัจจัยการผลิต")
table(doc, ["รายการ", "ชื่อวัสดุ", "ปริมาณ", "หน่วย", "ราคา/หน่วย", "มูลค่ารวม", "หลักฐาน"], [
    ["ปุ๋ยเคมี", "", "", "กก.", "", "", ""], ["ปุ๋ยอินทรีย์", "", "", "กก.", "", "", ""],
    ["สารกำจัดศัตรูพืช", "", "", "ลิตร/กก.", "", "", ""], ["สารชีวภัณฑ์", "", "", "ลิตร/กก.", "", "", ""],
    ["วัสดุปรับปรุงดิน", "", "", "กก.", "", "", ""],
], [1.15, 1.55, 0.72, 0.75, 0.85, 0.9, 0.9], font_size=8.2, first_col_left=True)
section(doc, "E2", "น้ำและพลังงาน")
table(doc, ["รายการ", "ปริมาณทั้งปี", "หน่วย", "ค่าใช้จ่ายบาท", "วิธีประมาณหรือหลักฐาน"], [
    ["น้ำเพื่อการเกษตร", "", "ลูกบาศก์เมตร", "", ""], ["ไฟฟ้า", "", "กิโลวัตต์ชั่วโมง", "", ""],
    ["น้ำมันดีเซล", "", "ลิตร", "", ""], ["น้ำมันเบนซิน", "", "ลิตร", "", ""],
], [1.45, 1.25, 1.2, 1.15, 2.0], font_size=8.6, first_col_left=True)
line(doc, "จำนวนครั้งที่ให้น้ำทั้งปี", 14, "ครั้ง  ชั่วโมงสูบน้ำรวม ________ ชั่วโมง")

# 10 F
page(doc, "แบบฟอร์ม F  การจัดการกิ่งและเศษชีวมวล", "บันทึกทุกครั้งที่นำกิ่งออกจากแปลงหรือเปลี่ยนวิธีจัดการ")
table(doc, ["วันที่", "แปลง", "วิธีจัดการ", "น้ำหนักสดกก.", "ความชื้นร้อยละ", "รหัสชุดเตา", "ค่าใช้จ่าย", "หลักฐาน"],
      [["", "", "", "", "", "", "", ""] for _ in range(9)],
      [0.68, 0.58, 1.25, 0.85, 0.9, 0.85, 0.85, 0.85], font_size=8.1)
options(doc, "รหัสวิธี", ["K เข้าเตาชีวมวล", "B เผากลางแจ้ง", "D ย่อยสลายในแปลง", "C ทำปุ๋ยหมัก", "O อื่น ๆ"])
section(doc, "F สรุป", "ยอดรวมตลอดปี")
line(doc, "น้ำหนักกิ่งสดทั้งหมด", 15, "กิโลกรัม  เข้าเตา ________ กก.  เผา ________ กก.")
line(doc, "ย่อยสลายหรือหมัก", 15, "กิโลกรัม  อื่น ๆ ________ กก.")
line(doc, "ปัญหาหรือข้อจำกัด", 58)

# 11 G
page(doc, "แบบฟอร์ม G  ระยะดอก ผล และการคาดการณ์", "บันทึกแยกตามแปลงเพื่อเชื่อมกับผลเก็บเกี่ยวจริง")
table(doc, ["วันที่", "แปลง", "ระยะ", "จำนวนผล", "คาดกก./ผล", "คาดวันเก็บ", "ปัญหา", "หลักฐาน"],
      [["", "", "", "", "", "", "", ""] for _ in range(9)],
      [0.68, 0.58, 1.15, 0.78, 0.8, 0.88, 1.25, 0.85], font_size=8.1)
options(doc, "ระยะ", ["ออกดอก", "ดอกบาน", "ติดผล", "พัฒนาผล", "ก่อนเก็บเกี่ยว"])
section(doc, "G ตรวจสอบ", "สภาพแปลงและความเสี่ยง")
options(doc, "พบปัญหา", ["โรค", "แมลง", "ผลร่วง", "ขาดน้ำ", "น้ำขัง", "ลมแรง", "ไม่พบ"])
line(doc, "แนวทางแก้ไขที่ทำจริง", 57)
line(doc, "ผู้สำรวจ", 30, "วันที่ตรวจสอบ ____ / ____ / ____")

# 12 H
page(doc, "แบบฟอร์ม H  การเก็บเกี่ยว", "บันทึกน้ำหนักจริงทุกครั้งและใช้วันที่เก็บผลสุดท้ายปิดรอบ")
table(doc, ["วันที่", "แปลง", "น้ำหนักกก.", "เกรด", "จำนวนผล", "ราคา/กก.", "ผู้ซื้อ", "หลักฐาน"],
      [["", "", "", "", "", "", "", ""] for _ in range(10)],
      [0.7, 0.58, 0.82, 0.7, 0.78, 0.85, 1.35, 0.85], font_size=8.2)
line(doc, "วันที่เก็บผลแรก", 22, "วันที่เก็บผลสุดท้าย ____ / ____ / ____")
line(doc, "น้ำหนักรวมทั้งรอบ", 16, "กิโลกรัม  จำนวนผลรวม ________ ผล")
line(doc, "ของเสียหรือผลตกเกรด", 16, "กิโลกรัม  สาเหตุ ____________________")
options(doc, "ตรวจหลักฐาน", ["ใบชั่ง", "ใบรับซื้อ", "ภาพถ่าย", "สมุดสวน", "สัมภาษณ์"])

# 13 I
page(doc, "แบบฟอร์ม I  การขายและรายได้", "บันทึกทั้งทุเรียนและผลิตภัณฑ์ชีวมวล")
table(doc, ["วันที่", "ผู้ซื้อหรือช่องทาง", "สินค้า", "ปริมาณ", "หน่วย", "ราคา/หน่วย", "มูลค่ารวม", "หลักฐาน"],
      [["", "", "", "", "", "", "", ""] for _ in range(9)],
      [0.68, 1.35, 1.15, 0.7, 0.6, 0.9, 0.92, 0.8], font_size=8.1)
options(doc, "สินค้า", ["ทุเรียน", "ถ่านชีวภาพ", "น้ำส้มควันไม้", "ผลิตภัณฑ์อื่น"])
options(doc, "ช่องทางหลัก", ["ล้ง", "พ่อค้าคนกลาง", "สหกรณ์", "ขายตรง", "ออนไลน์", "อื่น ๆ"])
line(doc, "รายได้รวมจากทุเรียน", 18, "บาท  รายได้ชีวมวล ________ บาท")
line(doc, "ยอดค้างรับ", 18, "บาท  วันที่คาดว่าจะได้รับ ____ / ____ / ____")

# 14 J
page(doc, "แบบฟอร์ม J  เทคโนโลยี ชีวมวล และคาร์บอน", "บันทึกการใช้งานจริง ต้นทุน ผลผลิต และกิจกรรมลดการปล่อย")
section(doc, "J1", "เทคโนโลยีและการเดินเตา")
table(doc, ["วันที่", "เครื่องหรือรหัสชุด", "วัตถุดิบกก.", "เวลาชม.", "ต้นทุนบาท", "ผลผลิต", "ปริมาณ", "หน่วย"],
      [["", "", "", "", "", "", "", ""] for _ in range(6)],
      [0.65, 1.2, 0.85, 0.7, 0.85, 1.15, 0.75, 0.75], font_size=8.2)
section(doc, "J2", "กิจกรรมลดคาร์บอน")
table(doc, ["วันที่", "กิจกรรม", "แปลง", "ปริมาณ", "หน่วย", "หลักฐาน"], [
    ["", "ใส่ถ่านชีวภาพลงดิน", "", "", "", ""], ["", "ลดปุ๋ยเคมีหรือใช้ปุ๋ยอินทรีย์", "", "", "", ""],
    ["", "ปลูกพืชคลุมดินหรือวนเกษตร", "", "", "", ""], ["", "ใช้พลังงานทดแทน", "", "", "", ""],
], [0.7, 2.35, 0.65, 0.85, 0.85, 1.55], font_size=8.4, first_col_left=False)

# 15 K
page(doc, "แบบฟอร์ม K  สรุปผลการจัดการสวนประจำปี", "ปิดรอบเมื่อบันทึกกิจกรรม เก็บเกี่ยว และการขายครบแล้ว")
section(doc, "K1", "ผลผลิตและเศรษฐกิจ")
table(doc, ["ตัวชี้วัด", "ค่า", "หน่วย", "วิธีคำนวณหรือตรวจสอบ"], [
    ["พื้นที่ให้ผลผลิต", "", "ไร่", "ทะเบียนแปลง"], ["ผลผลิตรวม", "", "กิโลกรัม", "รวมแบบ H"],
    ["ผลผลิตเฉลี่ย", "", "กก.ต่อไร่", "ผลผลิตรวม ÷ พื้นที่"], ["ต้นทุนรวม", "", "บาท", "รวมแบบ D E F J"],
    ["ต้นทุนต่อกิโลกรัม", "", "บาทต่อกก.", "ต้นทุนรวม ÷ ผลผลิต"], ["รายได้รวม", "", "บาท", "รวมแบบ I"],
    ["กำไรสุทธิประมาณการ", "", "บาท", "รายได้รวม - ต้นทุนรวม"], ["ราคาขายเฉลี่ย", "", "บาทต่อกก.", "รายได้ทุเรียน ÷ น้ำหนักขาย"],
], [2.25, 1.25, 1.35, 2.25], font_size=8.6, first_col_left=True)
section(doc, "K2", "บทเรียนและการเปลี่ยนแปลง")
line(doc, "สิ่งที่ทำได้ดีในปีนี้", 58)
line(doc, "ปัญหาสำคัญที่สุด", 60)
line(doc, "การเปลี่ยนแปลงจากปีก่อน", 52)
line(doc, "แผนปรับปรุงปีถัดไป", 58)
options(doc, "พร้อมปิดรอบ", ["ข้อมูลครบ", "รอตรวจสอบ", "ยังขาดข้อมูล"])

# 16 checklist and mapping
page(doc, "รายการตรวจสอบและแผนผังนำเข้าระบบ", "ใช้ก่อนส่งแบบให้ผู้บันทึกข้อมูล")
section(doc, "ตรวจความครบถ้วน", "ทำเครื่องหมายเมื่อผ่านการตรวจ")
for item in [
    "รหัสครัวเรือน สวน แปลง และรอบปีตรงกันทุกหน้า",
    "วันที่เรียงตามลำดับและอยู่ในรอบปีผลิตที่กำหนด",
    "ปริมาณ หน่วย ราคา ค่าใช้จ่าย และแรงงานระบุชัดเจน",
    "กิจกรรมรายเดือนเชื่อมกับหลักฐานหรือแหล่งข้อมูล",
    "ยอดเก็บเกี่ยว การขาย ต้นทุน และรายได้รวมตรวจทานแล้ว",
    "วันที่เก็บผลสุดท้ายระบุแล้วก่อนปิดรอบ",
    "ข้อมูลส่วนบุคคล พิกัด และภาพถ่ายมีความยินยอม",
]:
    options(doc, "", [item])
section(doc, "แผนผังนำเข้า", "นำข้อมูลจากแบบฟอร์มไปยังเมนูที่ตรงกัน")
table(doc, ["แบบ", "ข้อมูล", "เมนูในระบบ"], [
    ["A", "ครัวเรือน สวน แปลง", "ระบบครัวเรือน สวน และแปลง"],
    ["B", "รอบปีผลิต", "ฤดูการผลิต และรอบแปลง"],
    ["C D", "ปฏิทินและกิจกรรม", "Farm Activity Log"],
    ["E", "ปัจจัย น้ำ พลังงาน", "ข้อมูลฐานและต้นทุนการผลิต"],
    ["F", "กิ่งและเศษชีวมวล", "Branch Disposal"],
    ["G", "ระยะดอกและผล", "Phenology Records"],
    ["H", "เก็บเกี่ยว", "Harvest Records"],
    ["I", "การขาย", "Sales"],
    ["J", "เทคโนโลยี เตา คาร์บอน", "Kiln Batches Product Usage Carbon Activities"],
    ["K", "สรุปต้นทุน รายได้ ผลผลิต", "Production Cost และ Economic Impact"],
], [0.75, 2.25, 3.85], font_size=8.5, first_col_left=False)
line(doc, "ชื่อผู้บันทึกเข้าระบบ", 30, "วันที่ ____ / ____ / ____")
line(doc, "ชื่อผู้ตรวจสอบ", 34, "วันที่ ____ / ____ / ____")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
