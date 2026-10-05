from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "แบบสำรวจเกษตรกรชาวสวนทุเรียน_DRFI.docx"
FONT = "Noto Sans Thai"
GREEN = "1F6B35"
PALE_GREEN = "EAF4EC"
PALE_GRAY = "F5F6F7"
GRAY = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_font(run, size=11, bold=False, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text="", size=11, bold=False, align=None, space_after=4, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p


def add_labeled_line(doc, label, width=64, suffix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(label + "  ")
    set_font(r, bold=True)
    r = p.add_run("_" * width)
    set_font(r)
    if suffix:
        r = p.add_run("  " + suffix)
        set_font(r)
    return p


def add_options(doc, label, options, note=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(label + "  ")
    set_font(r, bold=True)
    for idx, option in enumerate(options):
        r = p.add_run(("   " if idx else "") + "☐ " + option)
        set_font(r)
    if note:
        r = p.add_run("  " + note)
        set_font(r)
    return p


def add_section_heading(doc, number, title, intro=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GREEN)
    p_pr.append(shd)
    r = p.add_run(f"ส่วนที่ {number}  {title}")
    set_font(r, size=14, bold=True)
    if intro:
        add_text(doc, intro, size=10, space_after=7)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        if widths:
            cell.width = Inches(widths[idx])
        set_cell_shading(cell, GREEN)
        set_cell_margins(cell, 100, 90, 100, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size=font_size, bold=True, color=WHITE)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            if widths:
                cells[idx].width = Inches(widths[idx])
            set_cell_margins(cells[idx], 110, 90, 110, 90)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2:
                set_cell_shading(cells[idx], PALE_GRAY)
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_font(r, size=18, bold=True)
    return p


def add_page_break(doc):
    doc.add_page_break()


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DRFIS  แบบสำรวจข้อมูลเกษตรกรชาวสวนทุเรียน  |  หน้า ")
    set_font(r, size=8, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)
add_footer(section)

styles = doc.styles
styles["Normal"].font.name = FONT
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Normal"].font.size = Pt(11)
styles["Title"].font.name = FONT
styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(22)
styles["Title"].font.bold = True

# หน้า 1
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
p.paragraph_format.space_after = Pt(8)
p.add_run("แบบสำรวจข้อมูลเกษตรกรชาวสวนทุเรียน")
add_text(doc, "สำหรับการบันทึกข้อมูลในระบบจัดการสวนทุเรียนและงานวิจัย DRFIS", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
add_text(doc, "พื้นที่ดำเนินงาน จังหวัดศรีสะเกษ", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=28)

add_section_heading(doc, "ก", "ข้อมูลการสัมภาษณ์")
add_labeled_line(doc, "รหัสแบบสำรวจ", 24, "รหัสครัวเรือนในระบบ ____________________")
add_labeled_line(doc, "วันที่สัมภาษณ์", 20, "เวลาเริ่ม ________ น.  เวลาสิ้นสุด ________ น.")
add_labeled_line(doc, "ชื่อผู้สัมภาษณ์", 35, "หน่วยงาน ____________________")
add_options(doc, "วิธีเก็บข้อมูล", ["สัมภาษณ์ ณ สวน", "สัมภาษณ์ ณ บ้าน", "โทรศัพท์", "อื่น ๆ __________"])

add_section_heading(doc, "ข", "คำชี้แจงและความยินยอม")
add_text(doc, "แบบสำรวจนี้ใช้เก็บข้อมูลเพื่อขึ้นทะเบียนครัวเรือน จัดการข้อมูลสวนและแปลง วิเคราะห์ต้นทุน ผลผลิต ตลาด เทคโนโลยี และผลกระทบด้านสิ่งแวดล้อม ข้อมูลส่วนบุคคลจะใช้ตามวัตถุประสงค์ของโครงการและจำกัดการเข้าถึงตามสิทธิ์ผู้ใช้งาน", size=10.5, space_after=8)
add_options(doc, "ผู้ให้ข้อมูลได้รับคำอธิบายและ", ["ยินยอมให้บันทึกข้อมูล", "ไม่ยินยอม"])
add_options(doc, "ยินยอมให้บันทึกพิกัดสวน", ["ยินยอม", "ไม่ยินยอม", "ไม่สามารถระบุพิกัด"])
add_options(doc, "ยินยอมให้ถ่ายภาพหลักฐานที่เกี่ยวข้อง", ["ยินยอม", "ไม่ยินยอม"])
add_labeled_line(doc, "ชื่อผู้ให้ข้อมูล", 35, "ลายมือชื่อ ____________________")
add_text(doc, "สำหรับเจ้าหน้าที่  หากไม่ยินยอม ให้ยุติการสัมภาษณ์และไม่บันทึกเลขประจำตัวประชาชนหรือพิกัด", size=9.5, bold=True, space_after=0)

# หน้า 2
add_page_break(doc)
add_page_title(doc, "ข้อมูลครัวเรือนและที่ตั้ง")
add_section_heading(doc, "1", "ข้อมูลผู้ให้ข้อมูลและครัวเรือน", "กรอกตามคำตอบของผู้ให้ข้อมูล ช่องที่ไม่ทราบให้ระบุว่าไม่ทราบ ห้ามคาดเดา")
add_labeled_line(doc, "ชื่อหัวหน้าครัวเรือน", 55)
add_labeled_line(doc, "เลขประจำตัวประชาชน 13 หลัก", 36, "ตรวจเฉพาะเมื่อได้รับความยินยอม")
add_labeled_line(doc, "เบอร์โทรศัพท์", 30, "ช่องทางติดต่ออื่น ____________________")
add_options(doc, "สถานะการเข้าร่วม", ["ใช้งานอยู่", "ระงับชั่วคราว", "ถอนตัว"])
add_labeled_line(doc, "วันที่เริ่มเข้าร่วมโครงการ", 26)
add_labeled_line(doc, "กลุ่มเกษตรกร", 48)
add_labeled_line(doc, "หมู่บ้าน", 30, "หมู่ที่ ________")
add_labeled_line(doc, "ตำบล", 24, "อำเภอ ____________________  จังหวัด ศรีสะเกษ")

add_section_heading(doc, "2", "ข้อมูลพื้นฐานครัวเรือน")
add_labeled_line(doc, "จำนวนสมาชิกในครัวเรือน", 12, "คน  สมาชิกที่ช่วยงานสวน ________ คน")
add_options(doc, "แรงงานหลักในสวน", ["ครัวเรือน", "จ้างประจำ", "จ้างรายวัน", "ผสมกัน"])
add_labeled_line(doc, "ประสบการณ์ปลูกทุเรียน", 12, "ปี")
add_options(doc, "มาตรฐานที่ได้รับ", ["GAP", "เกษตรอินทรีย์", "อื่น ๆ __________", "ยังไม่มี"])
add_labeled_line(doc, "ปัญหาสำคัญของครัวเรือนในการทำสวน", 62)
add_labeled_line(doc, "ความช่วยเหลือที่ต้องการ", 68)

# หน้า 3
add_page_break(doc)
add_page_title(doc, "ข้อมูลสวน")
add_section_heading(doc, "3", "ทะเบียนสวน", "ใช้หนึ่งชุดข้อมูลต่อหนึ่งสวน หากมีมากกว่า 3 สวน ให้แนบสำเนาหน้านี้เพิ่มเติม")
for n in range(1, 4):
    add_text(doc, f"สวนที่ {n}", size=12, bold=True, space_after=4, keep=True)
    add_labeled_line(doc, "ชื่อสวน", 28, "รหัสสวน ____________________")
    add_labeled_line(doc, "ที่อยู่หรือจุดสังเกต", 48)
    add_labeled_line(doc, "จังหวัด", 16, "อำเภอ ______________  ตำบล ______________")
    add_labeled_line(doc, "พื้นที่รวม", 12, "ไร่  แหล่งน้ำหลัก ____________________")
    add_labeled_line(doc, "พิกัดละติจูด", 17, "ลองจิจูด ____________________")
    add_options(doc, "สถานะสวน", ["ใช้งานอยู่", "หยุดใช้งาน"], note="จำนวนแปลง ________ แปลง")

# หน้า 4
add_page_break(doc)
add_page_title(doc, "ข้อมูลแปลงทุเรียน")
add_section_heading(doc, "4", "ทะเบียนแปลง", "บันทึกทุกแปลงในทุกสวน รหัสแปลงสามารถกำหนดภายหลังในระบบ")
headers = ["ลำดับ", "สวนที่", "รหัสแปลง", "พื้นที่\nไร่", "จำนวน\nต้น", "พันธุ์หลัก", "ปีปลูก", "ปีเริ่มให้ผล", "ระบบน้ำ", "สถานะ"]
rows = [[str(i), "", "", "", "", "", "", "", "", ""] for i in range(1, 9)]
add_table(doc, headers, rows, widths=[0.38, 0.42, 0.75, 0.52, 0.55, 0.9, 0.56, 0.7, 0.8, 0.62], font_size=8.2)
add_labeled_line(doc, "พันธุ์อื่นที่ปลูกร่วม", 55)
add_labeled_line(doc, "อายุการให้ผลผลิตเชิงเศรษฐกิจที่คาดไว้", 14, "ปี")
add_options(doc, "ลักษณะพื้นที่", ["ที่ราบ", "ลาดชัน", "เชิงเขา", "อื่น ๆ __________"])
add_options(doc, "การระบายน้ำ", ["ดี", "ปานกลาง", "มีน้ำขังเป็นครั้งคราว", "มีปัญหาน้ำขัง"])
add_labeled_line(doc, "หมายเหตุเกี่ยวกับแปลง", 64)

# หน้า 5
add_page_break(doc)
add_page_title(doc, "ข้อมูลฐานและการจัดการสวน")
add_section_heading(doc, "5", "ข้อมูลฐานของฤดูผลิต", "ใช้ข้อมูลของฤดูที่กำหนดเป็นฐานเปรียบเทียบก่อนเข้าร่วมกิจกรรมหรือใช้เทคโนโลยี")
add_labeled_line(doc, "ชื่อฤดูผลิตหรือปีการผลิต", 30, "ช่วงวันที่ ____________________ ถึง ____________________")
add_labeled_line(doc, "พื้นที่ให้ผลผลิต", 14, "ไร่  ผลผลิตรวม ____________________ กิโลกรัม")
add_labeled_line(doc, "ต้นทุนรวม", 18, "บาท  รายได้รวม ____________________ บาท")
add_labeled_line(doc, "ราคาขายเฉลี่ย", 15, "บาทต่อกิโลกรัม")
add_labeled_line(doc, "แนวทางจัดการสวนเดิม", 65)
add_labeled_line(doc, "การเปลี่ยนแปลงสำคัญจากปีก่อน", 58)

add_section_heading(doc, "6", "กิจกรรมการผลิตในฤดูปัจจุบัน", "บันทึกจากสมุดสวน ใบเสร็จ หรือคำให้สัมภาษณ์ และระบุแหล่งข้อมูลทุกครั้ง")
headers = ["วันที่", "แปลง", "กิจกรรม", "วัสดุหรือปัจจัย", "ปริมาณ", "หน่วย", "ค่าใช้จ่าย\nบาท", "แรงงาน\nชั่วโมง", "แหล่งข้อมูล"]
rows = [["", "", "", "", "", "", "", "", ""] for _ in range(8)]
add_table(doc, headers, rows, widths=[0.65, 0.55, 1.05, 1.15, 0.62, 0.55, 0.72, 0.72, 0.8], font_size=8.4)

# หน้า 6
add_page_break(doc)
add_page_title(doc, "ปัจจัยการผลิตและการจัดการกิ่ง")
add_section_heading(doc, "7", "ปริมาณใช้ต่อปี", "กรอกปริมาณจริง หากไม่ใช้ให้ใส่ 0 และระบุหน่วยให้ชัดเจน")
headers = ["รายการ", "ชื่อวัสดุหรือรายละเอียด", "ปริมาณต่อปี", "หน่วย", "ค่าใช้จ่ายบาท", "หลักฐานหรือหมายเหตุ"]
input_rows = [
    ["ปุ๋ยเคมี", "", "", "กก.", "", ""],
    ["ปุ๋ยอินทรีย์", "", "", "กก.", "", ""],
    ["สารกำจัดศัตรูพืช", "", "", "ลิตร/กก.", "", ""],
    ["ไฟฟ้า", "", "", "กิโลวัตต์ชั่วโมง", "", ""],
    ["น้ำมันดีเซล", "", "", "ลิตร", "", ""],
    ["น้ำมันเบนซิน", "", "", "ลิตร", "", ""],
    ["น้ำเพื่อการเกษตร", "", "", "ลบ.ม./ครั้ง", "", ""],
]
add_table(doc, headers, input_rows, widths=[1.05, 1.65, 0.9, 0.9, 0.95, 1.55], font_size=8.8)

add_section_heading(doc, "8", "การจัดการกิ่งและเศษชีวมวล")
add_options(doc, "วิธีจัดการหลัก", ["เข้าเตาชีวมวล", "เผากลางแจ้ง", "ย่อยสลายในแปลง", "ทำปุ๋ยหมัก", "อื่น ๆ ______"])
add_labeled_line(doc, "น้ำหนักสดโดยประมาณ", 15, "กิโลกรัมต่อปี  ความชื้น ________ เปอร์เซ็นต์")
add_labeled_line(doc, "ช่วงเดือนที่มีกิ่งมาก", 25)
add_labeled_line(doc, "ปัญหาในการจัดการกิ่ง", 58)

# หน้า 7
add_page_break(doc)
add_page_title(doc, "พัฒนาการผลและการเก็บเกี่ยว")
add_section_heading(doc, "9", "การติดตามระยะการเจริญ", "บันทึกแยกตามแปลงและฤดูผลิต")
headers = ["แปลง", "วันที่สำรวจ", "ระยะ", "จำนวนผลคงเหลือ", "น้ำหนักคาดการณ์\nกก.ต่อผล", "วันที่คาดเก็บเกี่ยว", "หมายเหตุ"]
rows = [["", "", "☐ ออกดอก  ☐ ดอกบาน  ☐ ติดผล  ☐ พัฒนาผล", "", "", "", ""] for _ in range(6)]
add_table(doc, headers, rows, widths=[0.65, 0.85, 1.85, 0.85, 0.9, 0.95, 0.95], font_size=8.4)

add_section_heading(doc, "10", "ผลผลิตและการเก็บเกี่ยว")
headers = ["วันที่เก็บ", "แปลง", "น้ำหนักจริง\nกก.", "เกรด", "ราคา\nบาทต่อกก.", "ผู้ซื้อ", "หลักฐาน"]
rows = [["", "", "", "", "", "", ""] for _ in range(7)]
add_table(doc, headers, rows, widths=[0.85, 0.65, 0.85, 0.72, 0.9, 1.5, 1.1], font_size=8.8)
add_labeled_line(doc, "วันที่เก็บผลสุดท้ายของรอบ", 25, "ใช้ปิดรอบการผลิตในระบบ")

# หน้า 8
add_page_break(doc)
add_page_title(doc, "การขายและตลาด")
add_section_heading(doc, "11", "รายการขาย", "บันทึกทั้งทุเรียนและผลิตภัณฑ์ชีวมวล โดยแยกแต่ละวันที่ขาย")
headers = ["วันที่", "ผู้ซื้อหรือช่องทาง", "สินค้า", "ปริมาณ", "หน่วย", "ราคา/หน่วย", "มูลค่ารวม", "ฤดูผลิต"]
rows = [["", "", "☐ ทุเรียน  ☐ ผลิตภัณฑ์ชีวมวล", "", "", "", "", ""] for _ in range(7)]
add_table(doc, headers, rows, widths=[0.7, 1.35, 1.3, 0.65, 0.55, 0.8, 0.85, 0.75], font_size=8.5)

add_section_heading(doc, "12", "สภาพตลาดและความต้องการ")
add_options(doc, "ช่องทางขายหลัก", ["ล้ง", "พ่อค้าคนกลาง", "สหกรณ์", "ขายตรง", "ออนไลน์", "อื่น ๆ ______"])
add_options(doc, "เงื่อนไขมาตรฐานจากผู้ซื้อ", ["GAP", "อินทรีย์", "ขนาด/เกรด", "ความสุก", "ไม่มี", "อื่น ๆ ______"])
add_labeled_line(doc, "ปริมาณที่ตลาดต้องการ", 18, "หน่วย __________  ความถี่ ____________________")
add_labeled_line(doc, "ราคาที่คาดหวัง", 18, "บาทต่อหน่วย")
add_options(doc, "มีข้อตกลงซื้อขาย", ["มีหนังสือ LOI/MOU", "ตกลงด้วยวาจา", "ยังไม่มี"])
add_labeled_line(doc, "ปัญหาด้านตลาด", 65)

# หน้า 9
add_page_break(doc)
add_page_title(doc, "เทคโนโลยีและผลิตภัณฑ์ชีวมวล")
add_section_heading(doc, "13", "เทคโนโลยีที่ได้รับหรือใช้งาน")
headers = ["ชื่อเทคโนโลยีหรือเครื่อง", "รหัสเครื่อง", "วันที่ได้รับ", "สถานะใช้งาน", "ผู้ดูแล", "ปัญหาหรือความต้องการ"]
rows = [["", "", "", "☐ ใช้ได้  ☐ ชำรุด  ☐ ส่งคืน", "", ""] for _ in range(4)]
add_table(doc, headers, rows, widths=[1.55, 0.85, 0.8, 1.2, 0.9, 1.7], font_size=8.8)

add_section_heading(doc, "14", "การเดินเตาและผลผลิตชีวมวล")
add_labeled_line(doc, "รหัสชุดการผลิต", 20, "วันที่เดินเตา ____________________")
add_labeled_line(doc, "วัตถุดิบเข้า", 16, "กิโลกรัม  เวลาผลิต __________ ชั่วโมง")
add_labeled_line(doc, "ค่าแรง", 14, "บาท  ค่าพลังงาน __________ บาท  ค่าอื่น ๆ __________ บาท")
headers = ["ผลิตภัณฑ์ที่ได้", "ปริมาณ", "หน่วย", "เกรด", "นำไปใช้หรือจำหน่ายอย่างไร"]
rows = [["", "", "", "", ""] for _ in range(4)]
add_table(doc, headers, rows, widths=[1.55, 0.9, 0.8, 0.75, 2.55], font_size=9)

# หน้า 10
add_page_break(doc)
add_page_title(doc, "กิจกรรมลดคาร์บอน")
add_section_heading(doc, "15", "กิจกรรมที่ดำเนินการ", "บันทึกเฉพาะสิ่งที่ทำจริง ระบุปริมาณและหน่วยเพื่อให้ระบบคำนวณคาร์บอนเทียบเท่า")
carbon_options = [
    "ใส่ถ่านชีวภาพลงดิน",
    "ใช้ปุ๋ยอินทรีย์ทดแทนปุ๋ยเคมี",
    "ลดการใช้ปุ๋ยเคมี",
    "ปลูกพืชคลุมดิน",
    "ปลูกไม้ยืนต้นร่วมแบบวนเกษตร",
    "ใช้พลังงานทดแทน",
]
for option in carbon_options:
    add_options(doc, "", [option], note="วันที่ __________  ปริมาณ __________  หน่วย __________")
add_labeled_line(doc, "รายละเอียดหรือหลักฐานประกอบ", 57)

add_section_heading(doc, "16", "การเปลี่ยนแปลงทางเศรษฐกิจ")
add_labeled_line(doc, "ต้นทุนที่ลดลง", 18, "บาทต่อฤดู")
add_labeled_line(doc, "รายได้ทุเรียนที่เพิ่มขึ้น", 18, "บาทต่อฤดู")
add_labeled_line(doc, "รายได้จากผลิตภัณฑ์ชีวมวล", 18, "บาทต่อฤดู")
add_labeled_line(doc, "ค่าใช้จ่ายเทคโนโลยีเพิ่มเติม", 18, "บาทต่อฤดู")
add_options(doc, "ผู้ให้ข้อมูลเห็นว่าผลประโยชน์สุทธิเพิ่มขึ้น", ["เพิ่มขึ้น", "ไม่เปลี่ยน", "ลดลง", "ยังประเมินไม่ได้"])

# หน้า 11
add_page_break(doc)
add_page_title(doc, "การยอมรับเทคโนโลยี")
add_section_heading(doc, "17", "ระดับการยอมรับเทคโนโลยี", "ประเมินแยกต่อเทคโนโลยี เลือกระดับสูงสุดที่ผู้ให้ข้อมูลทำได้จริง")
add_labeled_line(doc, "ชื่อเทคโนโลยี", 52)
alp_rows = [
    ["1", "รับรู้", "รู้จักหรืออธิบายประโยชน์ได้ แต่ยังไม่ทดลอง", "☐"],
    ["2", "ทดลอง", "เคยทดลองโดยมีผู้แนะนำหรือช่วยเหลือ", "☐"],
    ["3", "ใช้ได้เอง", "สามารถใช้งานได้ด้วยตนเองอย่างต่อเนื่อง", "☐"],
    ["4", "ปรับใช้และแก้ปัญหา", "ปรับเทคโนโลยีให้เหมาะกับสวนและแก้ปัญหาได้", "☐"],
    ["5", "ถ่ายทอดและขยายผล", "สาธิต สอน หรือช่วยให้ผู้อื่นนำไปใช้", "☐"],
]
add_table(doc, ["ระดับ", "ชื่อระดับ", "เกณฑ์พิจารณา", "เลือก"], alp_rows, widths=[0.6, 1.35, 4.65, 0.55], font_size=9.2)
add_labeled_line(doc, "เหตุผลหรือหลักฐานที่เลือกระดับนี้", 54)
add_labeled_line(doc, "อุปสรรคต่อการใช้เทคโนโลยี", 58)
add_labeled_line(doc, "การสนับสนุนที่ต้องการ", 63)

add_section_heading(doc, "18", "ความพึงพอใจต่อระบบและการสนับสนุน")
add_options(doc, "ความสะดวกในการบันทึกข้อมูล", ["มาก", "ค่อนข้างมาก", "ปานกลาง", "น้อย", "ยังไม่เคยใช้"])
add_options(doc, "ประโยชน์ที่ได้รับจากข้อมูล", ["มาก", "ค่อนข้างมาก", "ปานกลาง", "น้อย", "ยังประเมินไม่ได้"])
add_labeled_line(doc, "ข้อเสนอแนะต่อระบบหรือเจ้าหน้าที่", 58)

# หน้า 12
add_page_break(doc)
add_page_title(doc, "การประเมินสมรรถนะและการตรวจความครบถ้วน")
add_section_heading(doc, "19", "สมรรถนะนวัตกร", "ให้คะแนน 0 ถึง 100 เฉพาะกรณีประเมินนวัตกร ระบุรอบ T0 ก่อนพัฒนา T1 หลังฝึกหรือทดลอง หรือ T2 หลังใช้จริง")
add_options(doc, "รอบประเมิน", ["T0", "T1", "T2"], note="ประเภทผู้ประเมิน  ☐ ตนเอง  ☐ ผู้สังเกตการณ์")
competencies = [
    "หลักการเตาและกระบวนการผลิต", "การใช้ผลิตภัณฑ์ชีวมวล", "เตรียมวัตถุดิบและเดินเตา",
    "เก็บผลผลิตและความปลอดภัย", "แก้ปัญหาและปรับให้เหมาะกับสวน", "คำนวณต้นทุน ผลผลิต ราคา กำไร",
    "บรรจุภัณฑ์ แบรนด์ และช่องทางตลาด", "ความเข้าใจมาตรฐานสินค้า", "สาธิตและสอนผู้อื่น", "เป็นพี่เลี้ยงและขยายผล",
]
add_table(doc, ["ตัวชี้วัด", "คะแนนเต็ม", "คะแนนที่ได้", "หลักฐานหรือข้อสังเกต"], [[c, "100", "", ""] for c in competencies], widths=[3.2, 0.8, 0.9, 2.3], font_size=8.5)

add_section_heading(doc, "20", "รายการตรวจสอบก่อนนำเข้าระบบ")
checks = [
    "ตรวจชื่อและรหัสครัวเรือน ไม่สร้างข้อมูลซ้ำ",
    "เลือกหมู่บ้าน ตำบล อำเภอ และกลุ่มเกษตรกรแล้ว",
    "จำนวนสวนและแปลงตรงกับที่สัมภาษณ์",
    "หน่วยของพื้นที่ ปริมาณ ราคา และค่าใช้จ่ายชัดเจน",
    "ข้อมูลส่วนบุคคลและพิกัดมีความยินยอม",
    "แนบภาพหรือหลักฐานเฉพาะที่เกี่ยวข้อง",
    "ระบุข้อมูลที่ไม่ทราบโดยไม่คาดเดา",
]
for item in checks:
    add_options(doc, "", [item])
add_labeled_line(doc, "ชื่อผู้ตรวจสอบ", 34, "วันที่ ____________________")
add_labeled_line(doc, "หมายเหตุส่งต่อผู้บันทึกระบบ", 56)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
